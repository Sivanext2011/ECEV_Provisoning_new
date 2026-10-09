"""
`ecev batch` — CPM Batch: build / excel template+build / lifecycle.

Reuses backend batch logic directly:
  batch_mod.build_batch_file, _enrich_mandatory_chars, serialize_batch_file,
  _submit_raw, _fetch_result_stream, _extract_failure_reason, _build_multi,
  _combo_columns, _sample_row_for_combo, and ericsson_client.request for lifecycle.
"""
import io
import json
import click

from _common import (
    ericsson_client, catalog_svc, batch_mod, run, emit, table, err, load_json_file,
)


@click.group("batch")
def batch():
    """CPM Batch — build, Excel bulk, and job lifecycle."""


# ------------------------------------------------------------ build (single JSON file)
@batch.command("build")
@click.option("--input", "input_file", type=click.Path(exists=True), help="JSON build-options file.")
@click.option("--count", type=int, default=1, show_default=True)
@click.option("--po", help="Product offering externalId.")
@click.option("--adjustment", is_flag=True, help="Include balance adjustment entity.")
@click.option("--out", "out_file", type=click.Path(), help="Write the built batch file to this path.")
@click.pass_context
def batch_build(ctx, input_file, count, po, adjustment, out_file):
    """Build a batch file (party->customer->contract[->adjustment]) and print/save it."""
    body = load_json_file(input_file) if input_file else {}
    body.setdefault("count", count)
    if po:
        body["productOfferingExternalId"] = po
    if adjustment:
        body["adjustment"] = True
    bf = batch_mod.build_batch_file(body)
    bf = run(batch_mod._enrich_mandatory_chars(bf))
    if out_file:
        with open(out_file, "w", encoding="utf-8") as f:
            json.dump(bf, f, indent=2)
    emit(bf, ctx.obj["json"],
         lambda d: print(f"Built {len(bf.get('records', []))} record(s)." + (f" Saved to {out_file}" if out_file else "")))


# ------------------------------------------------------------ excel template
@batch.command("excel-template")
@click.option("--combos", "combos_file", type=click.Path(exists=True), help="JSON file: list of combo objects (see docs).")
@click.option("--sample-rows", type=int, default=1, show_default=True, help="Extra blank rows per template.")
@click.option("--out", "out_file", type=click.Path(), default="cpm_batch_bulk_template.xlsx", show_default=True)
@click.pass_context
def batch_excel_template(ctx, combos_file, sample_rows, out_file):
    """Generate a spec-driven Excel template (one sample row per combo)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    combos = load_json_file(combos_file) if combos_file else None
    if not combos:
        from _common import load_config
        d = load_config().get("defaults", {})
        combos = [{
            "comboName": "default",
            "partySpecExternalId": d.get("partySpecExternalId", ""),
            "customerSpecExternalId": d.get("customerSpecExternalId", ""),
            "billingAccountSpecExternalId": d.get("billingAccountSpecExternalId", ""),
            "billCycleSpecExternalId": d.get("billCycleSpecExternalId", ""),
            "contractSpecExternalId": d.get("contractSpecExternalId", ""),
            "productOfferingExternalId": d.get("basePlanProductOfferingExternalId", ""),
        }]
    specs = catalog_svc.get_catalog()

    wb = openpyxl.Workbook()
    hfont = Font(bold=True, color="FFFFFF")
    cfill = PatternFill("solid", fgColor="7C3AED")
    hfill = PatternFill("solid", fgColor="2563EB")

    # Combos sheet
    ws_c = wb.active
    ws_c.title = "Combos"
    combo_cols = ["comboName", "partySpecExternalId", "customerSpecExternalId",
                  "billingAccountSpecExternalId", "billCycleSpecExternalId",
                  "contractSpecExternalId", "productOfferingExternalId",
                  "contactMediumSpecExternalId", "resourceSpecs",
                  "includeBaRef", "includeBaRefRecurrence", "productPrice"]
    for ci, name in enumerate(combo_cols, 1):
        cc = ws_c.cell(row=1, column=ci, value=name); cc.font = hfont; cc.fill = cfill
    for ri, combo in enumerate(combos, 2):
        rs = combo.get("resourceSpecs") or []
        rs_enc = "|".join(f'{r.get("externalId","")}:{r.get("id","")}' for r in rs)
        pop = combo.get("productPrice") or []
        vals = [combo.get("comboName", f"combo{ri-1}"), combo.get("partySpecExternalId", ""),
                combo.get("customerSpecExternalId", ""), combo.get("billingAccountSpecExternalId", ""),
                combo.get("billCycleSpecExternalId", ""), combo.get("contractSpecExternalId", ""),
                combo.get("productOfferingExternalId", ""), combo.get("contactMediumSpecExternalId", ""),
                rs_enc, combo.get("includeBaRef", True), combo.get("includeBaRefRecurrence", True),
                json.dumps(pop) if pop else ""]
        for ci, v in enumerate(vals, 1):
            ws_c.cell(row=ri, column=ci, value=v)
    for ci in range(1, len(combo_cols) + 1):
        ws_c.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = 26

    # Entries sheet
    ws = wb.create_sheet("Entries")
    char_cols = batch_mod._combo_columns(specs, combos)
    all_cols = batch_mod._FIXED_COLS + [c[0] for c in char_cols]
    for ci, name in enumerate(all_cols, 1):
        cc = ws.cell(row=1, column=ci, value=name); cc.font = hfont; cc.fill = hfill
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = max(16, len(name) + 2)
    col_index = {name: i + 1 for i, name in enumerate(all_cols)}
    text_idxs = [col_index[n] for n in ("msisdn", "imsi") if n in col_index]
    r = 2
    for idx, combo in enumerate(combos, 1):
        sample = batch_mod._sample_row_for_combo(specs, combo, idx)
        for cn, v in sample.items():
            ci = col_index.get(cn)
            if ci:
                cell = ws.cell(row=r, column=ci, value=v)
                if ci in text_idxs:
                    cell.number_format = "@"
        r += 1
    first_combo = combos[0].get("comboName", "default")
    for _ in range(sample_rows):
        ws.cell(row=r, column=1, value=first_combo)
        for ci in text_idxs:
            ws.cell(row=r, column=ci).number_format = "@"
        r += 1

    # Instructions
    ws_i = wb.create_sheet("Instructions")
    for ri, line in enumerate([
        "CPM Batch — bulk entry template",
        "Combos sheet: named spec combinations (resourceSpecs = RS_MSISDN:id|RS_IMSI:id).",
        "Entries sheet: one row per subscriber; comboName picks the combo.",
        "party./customer./contract./cm./product.* columns = characteristic values.",
        "IMSI must be 15 digits. Upload with: ecev batch excel-build --input <file>.xlsx",
    ], 1):
        ws_i.cell(row=ri, column=1, value=line)
    ws_i.column_dimensions["A"].width = 90

    wb.save(out_file)
    emit({"template": out_file, "combos": [c.get("comboName") for c in combos], "columns": all_cols},
         ctx.obj["json"], lambda d: print(f"Wrote {out_file} ({len(combos)} combo(s), {len(all_cols)} columns)."))


# ------------------------------------------------------------ excel build (-> multi-record batch)
@batch.command("excel-build")
@click.option("--input", "xlsx_file", type=click.Path(exists=True), required=True, help="Filled .xlsx (Combos + Entries).")
@click.option("--out", "out_file", type=click.Path(), help="Write the built batch file JSON here.")
@click.option("--submit", is_flag=True, help="Submit (create) the job after building.")
@click.option("--start", is_flag=True, help="Also start the job (implies --submit).")
@click.pass_context
def batch_excel_build(ctx, xlsx_file, out_file, submit, start):
    """Parse a filled Excel into a multi-record batch; optionally submit/start."""
    import openpyxl
    with open(xlsx_file, "rb") as f:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    if "Combos" not in wb.sheetnames or "Entries" not in wb.sheetnames:
        err("workbook must contain 'Combos' and 'Entries' sheets"); raise SystemExit(1)

    # parse Combos
    c_rows = list(wb["Combos"].iter_rows(values_only=True))
    c_hdr = [str(h or "").strip() for h in c_rows[0]]
    combos_map = {}
    for rr in c_rows[1:]:
        if not any(v not in (None, "") for v in rr):
            continue
        rec = {c_hdr[i]: rr[i] for i in range(len(c_hdr)) if i < len(rr)}
        name = str(rec.get("comboName") or "").strip()
        if not name:
            continue
        rslist = []
        for part in str(rec.get("resourceSpecs") or "").split("|"):
            part = part.strip()
            if not part:
                continue
            ext, rid = (part.split(":", 1) + [""])[:2]
            rslist.append({"externalId": ext.strip(), "id": rid.strip()})
        def _b(v, d=True):
            return d if v in (None, "") else str(v).strip().lower() in ("true", "1", "yes", "y")
        pop = []
        if rec.get("productPrice") and str(rec["productPrice"]).strip():
            try:
                pop = json.loads(str(rec["productPrice"]))
            except Exception:
                pop = []
        combos_map[name] = {
            "comboName": name,
            "partySpecExternalId": rec.get("partySpecExternalId") or "",
            "customerSpecExternalId": rec.get("customerSpecExternalId") or "",
            "billingAccountSpecExternalId": rec.get("billingAccountSpecExternalId") or "",
            "billCycleSpecExternalId": rec.get("billCycleSpecExternalId") or "",
            "contractSpecExternalId": rec.get("contractSpecExternalId") or "",
            "productOfferingExternalId": rec.get("productOfferingExternalId") or "",
            "contactMediumSpecExternalId": rec.get("contactMediumSpecExternalId") or "",
            "resourceSpecs": rslist, "includeBaRef": _b(rec.get("includeBaRef")),
            "includeBaRefRecurrence": _b(rec.get("includeBaRefRecurrence")), "productPrice": pop,
        }
    if not combos_map:
        err("no combos in Combos sheet"); raise SystemExit(1)

    # parse Entries
    e_rows = list(wb["Entries"].iter_rows(values_only=True))
    e_hdr = [str(h or "").strip() for h in e_rows[0]]
    rows = []
    for rr in e_rows[1:]:
        if not any(v not in (None, "") for v in rr):
            continue
        row = {e_hdr[i]: rr[i] for i in range(len(e_hdr)) if i < len(rr)}
        data_fields = [k for k in row if k != "comboName"]
        if not any(str(row.get(k) or "").strip() for k in data_fields):
            continue
        rows.append(row)
    if not rows:
        err("no data rows in Entries"); raise SystemExit(1)

    bf = run(batch_mod._build_multi(rows, combos_map))
    if out_file:
        with open(out_file, "w", encoding="utf-8") as f:
            json.dump(bf, f, indent=2)

    result = {"recordCount": len(bf.get("records", [])), "combos": list(combos_map.keys())}
    if submit or start:
        created = run(batch_mod._submit_raw(bf))
        job_id = created.get("jobId") or created.get("id")
        result["jobId"] = job_id
        if start and job_id:
            run(ericsson_client.request("batch_start_job", path_params={"jobId": job_id}))
            result["started"] = True
    emit({**result, "batchFile": bf} if ctx.obj["json"] else result, ctx.obj["json"],
         lambda d: print(f"Built {result['recordCount']} record(s)."
                         + (f" jobId={result.get('jobId')}" if result.get("jobId") else "")
                         + (" (started)" if result.get("started") else "")
                         + (f" saved {out_file}" if out_file else "")))


# ------------------------------------------------------------ submit a prebuilt batch file
@batch.command("submit")
@click.option("--input", "bf_file", type=click.Path(exists=True), required=True, help="Batch file JSON ({header,records,trailer}).")
@click.option("--start", is_flag=True, help="Also start the job.")
@click.pass_context
def batch_submit(ctx, bf_file, start):
    """Submit a prebuilt batch file (create job)."""
    bf = load_json_file(bf_file)
    created = run(batch_mod._submit_raw(bf))
    job_id = created.get("jobId") or created.get("id")
    out = {"jobId": job_id}
    if start and job_id:
        run(ericsson_client.request("batch_start_job", path_params={"jobId": job_id}))
        out["started"] = True
    emit(out, ctx.obj["json"], lambda d: print(f"jobId={job_id}" + (" (started)" if out.get("started") else "")))


@batch.command("start")
@click.argument("job_id")
@click.pass_context
def batch_start(ctx, job_id):
    """Start a created job."""
    r = run(ericsson_client.request("batch_start_job", path_params={"jobId": job_id}))
    emit(r, ctx.obj["json"], lambda d: print(f"started {job_id}"))


@batch.command("status")
@click.argument("job_id")
@click.pass_context
def batch_status(ctx, job_id):
    """Get job status."""
    r = run(ericsson_client.request("batch_job_status", path_params={"jobId": job_id}))
    emit(r, ctx.obj["json"], lambda d: print(json.dumps(d, indent=2)))


@batch.command("result")
@click.argument("job_id")
@click.pass_context
def batch_result(ctx, job_id):
    """Get the full job result (parsed object stream)."""
    r = run(batch_mod._fetch_result_stream(job_id))
    emit(r, ctx.obj["json"])


@batch.command("failures")
@click.argument("job_id")
@click.pass_context
def batch_failures(ctx, job_id):
    """Show per-entity failure reasons for a job."""
    data = run(batch_mod._fetch_result_stream(job_id))
    objs = data.get("objects", [])
    total = ok = 0
    fails = []
    for o in objs:
        if "entities" not in o:
            continue
        for en in o.get("entities", []):
            total += 1
            if en.get("success"):
                ok += 1
                continue
            reason = batch_mod._extract_failure_reason(en.get("response") or "") or en.get("errorMessage") or ""
            fails.append({"recordNumber": o.get("recordNumber"), "entityNumber": en.get("entityNumber"),
                          "responseCode": en.get("responseCode"), "reason": reason})
    summary = {"jobId": job_id, "total": total, "success": ok, "failed": len(fails), "failures": fails}
    emit(summary, ctx.obj["json"], lambda d: (
        print(f"{ok} succeeded, {len(fails)} failed (of {total})."),
        table([[f["recordNumber"], f["entityNumber"], f["responseCode"], f["reason"]] for f in fails],
              ["rec", "entity", "code", "reason"]) if fails else print("No failures.")))


@batch.command("list")
@click.option("--sort", default="created_desc", show_default=True,
              type=click.Choice(["created_desc", "created_asc", "name_asc", "name_desc", "status_asc", "status_desc"]))
@click.option("--limit", type=int, default=20, show_default=True)
@click.pass_context
def batch_list(ctx, sort, limit):
    """List batch jobs (sorted)."""
    jobs = run(ericsson_client.request("batch_list_jobs"))
    if not isinstance(jobs, list):
        emit(jobs, ctx.obj["json"]); return
    import time as _t
    def epoch(j):
        s = (j.get("creationDate") or "").replace("\u202f", " ").strip()
        for fmt in ("%b %d, %Y, %I:%M:%S %p", "%b %d, %Y, %H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f%z", "%Y-%m-%dT%H:%M:%S"):
            try:
                return _t.mktime(_t.strptime(s, fmt))
            except Exception:
                continue
        return 0.0
    keymap = {
        "created_desc": (epoch, True), "created_asc": (epoch, False),
        "name_asc": (lambda j: (j.get("batchName") or "").lower(), False),
        "name_desc": (lambda j: (j.get("batchName") or "").lower(), True),
        "status_asc": (lambda j: (j.get("status") or "").lower(), False),
        "status_desc": (lambda j: (j.get("status") or "").lower(), True),
    }
    kf, rev = keymap.get(sort, keymap["created_desc"])
    jobs.sort(key=kf, reverse=rev)
    jobs = jobs[:limit]
    emit(jobs, ctx.obj["json"],
         lambda d: table([[j.get("jobId"), j.get("batchName"), j.get("status"), j.get("creationDate")] for j in jobs],
                         ["jobId", "name", "status", "created"]))


@batch.command("delete")
@click.argument("job_id")
@click.confirmation_option(prompt="Delete this batch job?")
@click.pass_context
def batch_delete(ctx, job_id):
    """Delete a batch job."""
    r = run(ericsson_client.request("batch_delete_job", path_params={"jobId": job_id}))
    emit(r, ctx.obj["json"], lambda d: print(f"deleted {job_id}"))
